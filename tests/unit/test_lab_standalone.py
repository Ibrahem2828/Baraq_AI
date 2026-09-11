from __future__ import annotations

import io
from datetime import date, timedelta
from pathlib import Path

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.api.v1.lab import router as lab_router
from app.application.standalone import BaraqAIApplication
from app.core.config import Settings
from app.core.errors import ValidationFailure
from app.lab.jobs import LocalJobManager
from app.lab.providers import build_lab_provider
from app.lab.retrieval import LocalLexicalRetriever
from app.lab.storage import LabStorage
from app.models.enums import TaskType
from app.providers.mock_provider import MockProvider


def _settings(tmp_path: Path) -> Settings:
    return Settings(lab_storage_dir=tmp_path / ".baraq_lab", provider_mode="mock")


def _storage(tmp_path: Path) -> LabStorage:
    return LabStorage(_settings(tmp_path))


def _text_pdf() -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        (
            b"<< /Length 47 >>\nstream\nBT /F1 14 Tf 72 720 Td "
            b"(Physics source evidence) Tj ET\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    chunks = [b"%PDF-1.4\n"]
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(sum(len(chunk) for chunk in chunks))
        chunks.append(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_offset = sum(len(chunk) for chunk in chunks)
    xref = [b"xref\n0 6\n0000000000 65535 f \n"]
    xref.extend(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    chunks.extend(xref)
    chunks.append(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode())
    return b"".join(chunks)


def test_lab_persists_txt_and_limits_retrieval_to_selected_source(tmp_path: Path) -> None:
    storage = _storage(tmp_path)
    workspace = storage.ensure_workspace()
    physics = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename="physics.txt",
        mime_type="text/plain",
        content=b"Newton force equals mass times acceleration.",
    )
    history = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename="history.txt",
        mime_type="text/plain",
        content=b"The treaty was signed after the conference.",
    )
    chunks = storage.list_chunks(
        workspace_id=workspace.workspace_id, source_ids=[physics.source_id]
    )
    evidence = LocalLexicalRetriever().retrieve(chunks=chunks, query="force acceleration", limit=3)

    assert physics.status == "ready"
    assert all(item.chunk.source_id == physics.source_id for item in evidence)
    assert history.source_id not in {item.chunk.source_id for item in evidence}


@pytest.mark.parametrize(
    "filename,mime_type,content", [("study.pdf", "application/pdf", _text_pdf())]
)
def test_lab_processes_readable_pdf(
    tmp_path: Path, filename: str, mime_type: str, content: bytes
) -> None:
    storage = _storage(tmp_path)
    workspace = storage.ensure_workspace()
    source = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename=filename,
        mime_type=mime_type,
        content=content,
    )
    chunks = storage.list_chunks(workspace_id=workspace.workspace_id, source_ids=[source.source_id])
    assert source.metadata["extractor"] == "pypdf"
    assert "Physics" in chunks[0].text


def test_lab_processes_docx_and_pptx_tables(tmp_path: Path) -> None:
    storage = _storage(tmp_path)
    workspace = storage.ensure_workspace()
    document = Document()
    document.add_heading("Biology", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "cell"
    table.cell(0, 1).text = "membrane"
    docx = io.BytesIO()
    document.save(docx)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Chemistry"
    ppt_table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    ppt_table.cell(0, 0).text = "acid"
    ppt_table.cell(0, 1).text = "base"
    pptx = io.BytesIO()
    presentation.save(pptx)

    docx_source = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename="biology.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=docx.getvalue(),
    )
    pptx_source = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename="chemistry.pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        content=pptx.getvalue(),
    )
    docx_text = " ".join(
        item.text
        for item in storage.list_chunks(
            workspace_id=workspace.workspace_id, source_ids=[docx_source.source_id]
        )
    )
    pptx_text = " ".join(
        item.text
        for item in storage.list_chunks(
            workspace_id=workspace.workspace_id, source_ids=[pptx_source.source_id]
        )
    )
    assert "membrane" in docx_text
    assert "acid" in pptx_text


def test_lab_rejects_mime_mismatch(tmp_path: Path) -> None:
    storage = _storage(tmp_path)
    with pytest.raises(ValidationFailure, match="MIME"):
        storage.store_upload(
            workspace_id=storage.ensure_workspace().workspace_id,
            filename="disguised.pdf",
            mime_type="text/plain",
            content=b"not a pdf",
        )


@pytest.mark.asyncio
async def test_khota_uses_deterministic_local_scheduler_without_provider(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    storage = LabStorage(settings)
    workspace = storage.ensure_workspace()
    application = BaraqAIApplication(settings=settings, storage=storage, provider=MockProvider())
    start = date.today()
    output = await application.run_task(
        workspace_id=workspace.workspace_id,
        task_type=TaskType.KHOTA_GENERATE_PLAN,
        input={
            "subject_ids": ["math"],
            "start_date": start.isoformat(),
            "end_date": (start + timedelta(days=1)).isoformat(),
            "daily_available_minutes": 45,
            "preferred_session_minutes": 30,
        },
    )
    assert output.provider["name"] == "local_scheduler"
    assert output.result["plan_days"]


def test_lab_browser_and_health_are_available_in_mock_mode(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    storage = LabStorage(settings)
    application = BaraqAIApplication(settings=settings, storage=storage, provider=MockProvider())
    app = FastAPI()
    app.state.lab_storage = storage
    app.state.lab_application = application
    app.state.lab_jobs = LocalJobManager(storage=storage, application=application)
    app.include_router(lab_router)

    with TestClient(app) as client:
        home = client.get("/lab", follow_redirects=True)
        health = client.get("/lab/health")
        assert home.status_code == 200
        assert "مختبر برّاق" in home.text
        body = health.json()
        assert body["provider_mode"] == "mock"
        assert body["configured"] is True


def test_lab_health_reports_fallback_when_live_mode_has_no_credentials(tmp_path: Path) -> None:
    # Explicit empty credentials so this stays hermetic even when a real
    # OPENAI_PRIMARY_API_KEY is present in the developer's local .env -- this
    # test asserts the *no credentials* fallback path specifically, not
    # whatever key happens to be on disk.
    settings = Settings(
        lab_storage_dir=tmp_path / ".baraq_lab",
        provider_mode="live",
        openai_primary_api_key="",
        gemini_primary_api_key="",
    )
    storage = LabStorage(settings)
    provider = build_lab_provider(settings)
    assert isinstance(provider, MockProvider)
    application = BaraqAIApplication(settings=settings, storage=storage, provider=provider)
    app = FastAPI()
    app.state.lab_storage = storage
    app.state.lab_application = application
    app.state.lab_jobs = LocalJobManager(storage=storage, application=application)
    app.include_router(lab_router)

    with TestClient(app) as client:
        health = client.get("/lab/health")
        body = health.json()
        assert body["configured"] is False
        assert body["message"] is not None


@pytest.mark.asyncio
async def test_sada_rejects_a_non_audio_source_before_stt_or_provider(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    storage = LabStorage(settings)
    workspace = storage.ensure_workspace()
    text_source = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename="notes.txt",
        mime_type="text/plain",
        content=b"Study notes",
    )
    app = BaraqAIApplication(settings=settings, storage=storage, provider=MockProvider())
    with pytest.raises(ValidationFailure, match="audio"):
        await app.run_task(
            workspace_id=workspace.workspace_id,
            task_type=TaskType.SADA_TRANSCRIBE_AUDIO,
            input={"source_id": text_source.source_id},
        )


@pytest.mark.asyncio
async def test_fahes_runs_end_to_end_against_mock_provider(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    storage = LabStorage(settings)
    workspace = storage.ensure_workspace()
    # Content chosen to lexically overlap with MockProvider's fixed Arabic
    # quiz fixture (app/providers/mock_provider.py) so the real
    # ClaimEvidenceValidator grounding check -- not a stub -- can pass.
    source = storage.store_upload(
        workspace_id=workspace.workspace_id,
        filename="geography.txt",
        mime_type="text/plain",
        content="المصدر يذكر الرياض كمثال توضيحي رقم واحد على عاصمة عربية.".encode(),
    )
    application = BaraqAIApplication(settings=settings, storage=storage, provider=MockProvider())
    output = await application.run_task(
        workspace_id=workspace.workspace_id,
        task_type=TaskType.FAHES_GENERATE_QUIZ,
        input={"source_ids": [source.source_id], "question_count": 3, "language": "ar"},
    )
    assert output.validation["status"] == "valid"
    assert output.provider["model"]
    assert output.result["questions"]
