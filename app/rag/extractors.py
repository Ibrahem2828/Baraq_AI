from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import ClassVar

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from app.core.errors import ValidationFailure


@dataclass(frozen=True, slots=True)
class ExtractedPage:
    page_number: int | None
    text: str
    section_title: str | None = None


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    pages: list[ExtractedPage]
    metadata: dict[str, str | int | float | bool | None]

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text.strip())


class DocumentExtractor:
    TEXT_TYPES: ClassVar[set[str]] = {
        "text/plain",
        "text/markdown",
        "text/csv",
        "application/json",
    }

    def extract(self, *, filename: str, mime_type: str, content: bytes) -> ExtractedDocument:
        suffix = Path(filename).suffix.lower()
        if mime_type in self.TEXT_TYPES or suffix in {".txt", ".md", ".csv", ".json"}:
            return self._extract_text(content)
        if mime_type == "application/pdf" or suffix == ".pdf":
            return self._extract_pdf(content)
        if (
            mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or suffix == ".docx"
        ):
            return self._extract_docx(content)
        if (
            mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            or suffix == ".pptx"
        ):
            return self._extract_pptx(content)
        raise ValidationFailure(
            "Unsupported source format",
            code="unsupported_source_format",
            details={"filename": filename, "mime_type": mime_type},
        )

    @staticmethod
    def _decode(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValidationFailure("Unable to decode text source", code="text_decode_failed")

    def _extract_text(self, content: bytes) -> ExtractedDocument:
        text = self._decode(content).replace("\x00", "").strip()
        if not text:
            raise ValidationFailure("The source contains no readable text", code="empty_source")
        return ExtractedDocument([ExtractedPage(None, text)], {"extractor": "plain_text"})

    def _extract_pdf(self, content: bytes) -> ExtractedDocument:
        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as exc:
            raise ValidationFailure("Unable to read PDF", code="pdf_read_failed") from exc
        pages: list[ExtractedPage] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").replace("\x00", "").strip()
            if text:
                pages.append(ExtractedPage(index, text))
        if not pages:
            raise ValidationFailure(
                "PDF contains no extractable text; OCR is required",
                code="pdf_ocr_required",
            )
        return ExtractedDocument(
            pages,
            {"extractor": "pypdf", "page_count": len(reader.pages)},
        )

    def _extract_docx(self, content: bytes) -> ExtractedDocument:
        try:
            document = Document(io.BytesIO(content))
        except Exception as exc:
            raise ValidationFailure("Unable to read DOCX", code="docx_read_failed") from exc
        blocks: list[str] = []
        current_heading: str | None = None
        pages: list[ExtractedPage] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                if blocks:
                    pages.append(ExtractedPage(None, "\n".join(blocks), current_heading))
                    blocks = []
                current_heading = text
            else:
                blocks.append(text)
        if blocks:
            pages.append(ExtractedPage(None, "\n".join(blocks), current_heading))
        # `python-docx` exposes tables separately from paragraphs.  Append a
        # compact, searchable representation instead of silently discarding
        # source facts that happen to be tabular.
        table_blocks = [self._docx_table_text(table) for table in document.tables]
        table_blocks = [block for block in table_blocks if block]
        if table_blocks:
            pages.append(ExtractedPage(None, "\n\n".join(table_blocks), "Tables"))
        if not pages:
            raise ValidationFailure("DOCX contains no readable text", code="empty_source")
        return ExtractedDocument(pages, {"extractor": "python-docx"})

    @staticmethod
    def _docx_table_text(table: object) -> str:
        rows = getattr(table, "rows", [])
        rendered: list[str] = []
        for row in rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                rendered.append(" | ".join(values))
        return "\n".join(rendered)

    def _extract_pptx(self, content: bytes) -> ExtractedDocument:
        try:
            presentation = Presentation(io.BytesIO(content))
        except Exception as exc:
            raise ValidationFailure("Unable to read PPTX", code="pptx_read_failed") from exc
        pages: list[ExtractedPage] = []
        for slide_no, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            title: str | None = None
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    rows = [
                        " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                        for row in shape.table.rows
                    ]
                    texts.extend(row for row in rows if row.strip(" |"))
                if not hasattr(shape, "text"):
                    continue
                value = str(shape.text).strip()
                if not value:
                    continue
                if title is None and getattr(shape, "is_placeholder", False):
                    title = value
                texts.append(value)
            if texts:
                pages.append(ExtractedPage(slide_no, "\n".join(texts), title))
        if not pages:
            raise ValidationFailure("PPTX contains no readable text", code="empty_source")
        return ExtractedDocument(
            pages,
            {"extractor": "python-pptx", "slide_count": len(presentation.slides)},
        )
